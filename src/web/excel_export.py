"""Generate Excel (.xlsx) reports from RunResult data."""

from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.models.test_result import RunResult

_GREEN = PatternFill("solid", fgColor="C6EFCE")
_RED = PatternFill("solid", fgColor="FFC7CE")
_YELLOW = PatternFill("solid", fgColor="FFEB9C")
_HEADER_FILL = PatternFill("solid", fgColor="2B4C7E")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _header_row(ws: Any, values: list[str]) -> None:
    ws.append(values)
    for col_idx, _ in enumerate(values, start=1):
        cell = ws.cell(row=ws.max_row, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _autofit(ws: Any) -> None:
    for col_cells in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max_len + 4, 60)


def _result_fill(result: str) -> PatternFill | None:
    mapping = {"pass": _GREEN, "fail": _RED, "error": _RED, "skip": _YELLOW}
    return mapping.get(result.lower())


def generate_xlsx(run_result: RunResult) -> bytes:
    wb = Workbook()

    # ── Sheet 1: Summary ─────────────────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = "Summary"
    _header_row(ws_summary, ["Field", "Value"])
    rows = [
        ("Run ID", run_result.run_id),
        ("Target URL", run_result.target_url),
        ("Started At", run_result.started_at),
        ("Completed At", run_result.completed_at),
        ("Duration (s)", round(run_result.duration_seconds, 1)),
        ("Total Tests", run_result.total_tests),
        ("Passed", run_result.passed),
        ("Failed", run_result.failed),
        ("Skipped", run_result.skipped),
        ("Errors", run_result.errors),
        ("AI Summary", run_result.ai_summary or "—"),
    ]
    for key, val in rows:
        ws_summary.append([key, val])
    _autofit(ws_summary)

    # ── Sheet 2: Test Cases ───────────────────────────────────────────────────
    ws_tests = wb.create_sheet("Test Cases")
    _header_row(
        ws_tests,
        ["Test ID", "Test Name", "Category", "Result", "Duration (s)", "Failure Reason"],
    )
    for t in run_result.test_results:
        row = [
            t.test_id,
            t.test_name,
            t.category,
            t.result.upper(),
            round(t.duration_seconds, 2),
            t.failure_reason or "",
        ]
        ws_tests.append(row)
        fill = _result_fill(t.result)
        if fill:
            result_cell = ws_tests.cell(row=ws_tests.max_row, column=4)
            result_cell.fill = fill
    _autofit(ws_tests)

    # ── Sheet 3: Errors ───────────────────────────────────────────────────────
    ws_errors = wb.create_sheet("Errors")
    _header_row(
        ws_errors,
        ["Test ID", "Test Name", "Category", "Result", "Failure Reason", "Failed Step"],
    )
    for t in run_result.test_results:
        if t.result.lower() not in ("fail", "error"):
            continue
        failed_step = ""
        for step in t.step_results:
            if step.status in ("fail", "error"):
                failed_step = step.description or step.action_type
                break
        ws_errors.append([
            t.test_id,
            t.test_name,
            t.category,
            t.result.upper(),
            t.failure_reason or "",
            failed_step,
        ])
        for col_idx in range(1, 7):
            ws_errors.cell(row=ws_errors.max_row, column=col_idx).fill = _RED
    _autofit(ws_errors)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
